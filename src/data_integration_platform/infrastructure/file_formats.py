"""受管文件格式的模式推断与受限采样（T-16 / FR-09）。

九种格式共用同一契约：字段清单 + 行数（**引擎能给就给，给不了留空**）+ 受限样本。
本模块是纯函数库，不碰磁盘也不碰网络：字节由调用方按各自边界读入
（本地资产 25 MiB、S3 前缀范围采样 8 MiB、REST 4 MiB），因此「有界读取」不可能被绕过。

原始取值只在内存里流转；脱敏统一由 `ProfileService` 完成。
"""

import csv
import io
import json
import re
from collections.abc import Iterable, Mapping, Sequence
from dataclasses import dataclass, field
from typing import Any

from defusedxml import ElementTree as SafeElementTree

from data_integration_platform.application.profiling import ProbedColumn
from data_integration_platform.domain.enums import FileAssetFormat
from data_integration_platform.domain.errors import DomainError
from data_integration_platform.infrastructure.probe_support import cell_to_plain

# 只映射到 PRD §8.4 声明的九种格式；不认识的后缀明确报错，不猜、不静默降级。
_SUFFIX_FORMATS: Mapping[str, FileAssetFormat] = {
    ".csv": FileAssetFormat.CSV,
    ".tsv": FileAssetFormat.CSV,
    ".json": FileAssetFormat.JSON,
    ".jsonl": FileAssetFormat.JSONL,
    ".ndjson": FileAssetFormat.JSONL,
    ".xlsx": FileAssetFormat.XLSX,
    ".xlsm": FileAssetFormat.XLSX,
    ".parquet": FileAssetFormat.PARQUET,
    ".orc": FileAssetFormat.ORC,
    ".xml": FileAssetFormat.XML,
    ".md": FileAssetFormat.MARKDOWN,
    ".markdown": FileAssetFormat.MARKDOWN,
    ".pdf": FileAssetFormat.PDF,
}

_FILE_ASSET_FORMATS = frozenset(_SUFFIX_FORMATS.values())
_TEXT_ENCODINGS = ("utf-8-sig", "utf-8", "gb18030")
_CSV_DELIMITERS = (",", ";", "\t", "|")
_MARKDOWN_CELL_SPLIT = re.compile(r"\|")
_MARKDOWN_ALIGNMENT_ROW = re.compile(r"^\s*\|?[\s:|-]+\|?\s*$")


@dataclass(frozen=True, slots=True)
class FileInspection:
    """一次文件模式推断的结果。`row_count` 为 `None` 表示该格式未提供行数。"""

    columns: tuple[ProbedColumn, ...]
    row_count: int | None
    sample_rows: tuple[Mapping[str, Any], ...] = ()
    statistics: Mapping[str, Any] = field(default_factory=dict)


def detect_format(name: str) -> FileAssetFormat:
    """按扩展名判定格式；不支持的后缀直接拒绝，不做内容嗅探式猜测。"""
    lowered = name.strip().casefold()
    for suffix, file_format in _SUFFIX_FORMATS.items():
        if lowered.endswith(suffix):
            return file_format
    raise DomainError(f"不支持的文件后缀：{name}")


def supported_formats() -> frozenset[FileAssetFormat]:
    return _FILE_ASSET_FORMATS


# 元数据位于**文件尾部**或需要完整容器才能解析的格式：前缀采样对它们不成立。
# Parquet / ORC 的 schema 在 footer，XLSX 是 zip、PDF 的 trailer 在末尾。
_FOOTER_BASED_FORMATS = frozenset(
    {
        FileAssetFormat.PARQUET,
        FileAssetFormat.ORC,
        FileAssetFormat.XLSX,
        FileAssetFormat.PDF,
    }
)


def requires_complete_file(file_format: FileAssetFormat) -> bool:
    """该格式能否仅凭前缀字节推断结构；不能则调用方必须整份读取或明确报错。"""
    return file_format in _FOOTER_BASED_FORMATS


def inspect_file(
    data: bytes,
    *,
    file_format: FileAssetFormat,
    sample_limit: int,
    encoding: str | None = None,
    complete: bool = True,
) -> FileInspection:
    """按格式推断模式；所有分支的读取都基于已读入的 `data`，不再触达外部。

    `complete=False` 表示 `data` 只是文件的前缀（对象存储/WebHDFS 的受限读取）。
    此时**行数一律留空** —— 用部分字节数出来的行数不是文件真实行数，
    把它当事实汇报会比不给更糟。
    """
    if not data:
        raise DomainError("文件内容为空")
    if not complete and requires_complete_file(file_format):
        raise DomainError(
            f"{file_format.value} 的结构信息位于文件尾部，无法通过前缀采样推断；"
            "请使用不超过采样上限的对象，或改用文本类格式"
        )
    handler = _HANDLERS.get(file_format)
    if handler is None:
        raise DomainError(f"{file_format.value} 格式探查器尚未接入")
    inspection = handler(data, sample_limit, encoding)
    if complete:
        return inspection
    return FileInspection(
        columns=inspection.columns,
        row_count=None,
        sample_rows=inspection.sample_rows,
        statistics={**inspection.statistics, "partial_read": True, "sampled_bytes": len(data)},
    )


def _decode(data: bytes, encoding: str | None) -> tuple[str, str]:
    """解码文本；返回（文本, 实际使用的编码），全部失败时用替换字符兜底而不是崩掉。"""
    candidates = (encoding,) if encoding else _TEXT_ENCODINGS
    for candidate in candidates:
        try:
            return data.decode(candidate), candidate
        except (UnicodeDecodeError, LookupError):
            continue
    return data.decode("utf-8", errors="replace"), "utf-8/replace"


def _columns(pairs: Iterable[tuple[str, str]]) -> tuple[ProbedColumn, ...]:
    """按出现顺序生成字段清单；重名后者加序号，避免样本字典互相覆盖。"""
    seen: dict[str, int] = {}
    columns: list[ProbedColumn] = []
    for ordinal, (name, data_type) in enumerate(pairs, start=1):
        unique = name
        if name in seen:
            seen[name] += 1
            unique = f"{name}_{seen[name]}"
        else:
            seen[name] = 1
        columns.append(ProbedColumn(name=unique, data_type=data_type, ordinal=ordinal))
    return tuple(columns)


def _content_only(note: str, **statistics: Any) -> FileInspection:
    """非表格型格式的保守结果：只声明一个文本字段，**不产出样本**。

    PDF 与纯文本 Markdown 的正文可能整段是个人信息，值级兜底检测对长文本不可靠，
    因此这里选择不产出样本 —— 宁可不给，也不冒泄露风险。
    """
    return FileInspection(
        columns=(ProbedColumn(name="content", data_type="TEXT", ordinal=1),),
        row_count=None,
        sample_rows=(),
        statistics={"representation": note, **statistics},
    )


# ------------------------------------------------------------------ 定界文本
def _sniff_delimiter(first_line: str) -> str:
    counts = {delimiter: first_line.count(delimiter) for delimiter in _CSV_DELIMITERS}
    best = max(counts, key=lambda item: counts[item])
    return best if counts[best] > 0 else ","


def _inspect_csv(data: bytes, sample_limit: int, encoding: str | None) -> FileInspection:
    text, used_encoding = _decode(data, encoding)
    reader = csv.reader(
        io.StringIO(text, newline=""), delimiter=_sniff_delimiter(text.split("\n", 1)[0])
    )
    rows = [row for row in reader if row]
    if not rows:
        raise DomainError("CSV 文件没有可解析的行")
    header = [cell.strip() or f"column_{index}" for index, cell in enumerate(rows[0])]
    body = rows[1:]
    return FileInspection(
        columns=_columns((name, "STRING") for name in header),
        row_count=len(body),
        sample_rows=tuple(
            {
                header[index]: cell_to_plain(cell)
                for index, cell in enumerate(row)
                if index < len(header)
            }
            for row in body[:sample_limit]
        ),
        statistics={"encoding": used_encoding, "delimiter": reader.dialect.delimiter},
    )


# ------------------------------------------------------------------ JSON 系列
def _json_columns(records: Sequence[Mapping[str, Any]]) -> tuple[ProbedColumn, ...]:
    ordered: dict[str, str] = {}
    for record in records:
        for key, value in record.items():
            ordered.setdefault(key, _json_type(value))
    return _columns(ordered.items())


def _json_type(value: Any) -> str:
    if value is None:
        return "NULL"
    if isinstance(value, bool):
        return "BOOLEAN"
    if isinstance(value, int):
        return "INTEGER"
    if isinstance(value, float):
        return "DOUBLE"
    if isinstance(value, str):
        return "STRING"
    if isinstance(value, list | tuple):
        return "ARRAY"
    if isinstance(value, dict):
        return "OBJECT"
    return "STRING"


def _inspect_json(data: bytes, sample_limit: int, encoding: str | None) -> FileInspection:
    text, used_encoding = _decode(data, encoding)
    try:
        parsed = json.loads(text)
    except json.JSONDecodeError as exc:
        raise DomainError(f"JSON 解析失败：{exc.msg}（第 {exc.lineno} 行）") from exc
    if isinstance(parsed, list):
        records = [item for item in parsed if isinstance(item, dict)]
        if not records and parsed:
            # 标量数组无法表达为字段清单，明确报错而不是造一个假字段。
            raise DomainError("JSON 顶层数组的元素不是对象，无法推断字段清单")
        return FileInspection(
            columns=_json_columns(records),
            row_count=len(parsed),
            sample_rows=tuple(dict(item) for item in records[:sample_limit]),
            statistics={"encoding": used_encoding, "root": "array"},
        )
    if isinstance(parsed, dict):
        return FileInspection(
            columns=_columns((key, _json_type(value)) for key, value in parsed.items()),
            row_count=1,
            sample_rows=(dict(parsed),),
            statistics={"encoding": used_encoding, "root": "object"},
        )
    raise DomainError("JSON 顶层既不是对象也不是数组，无法推断字段清单")


def _inspect_jsonl(data: bytes, sample_limit: int, encoding: str | None) -> FileInspection:
    text, used_encoding = _decode(data, encoding)
    records: list[dict[str, Any]] = []
    total = 0
    for line_number, line in enumerate(text.splitlines(), start=1):
        stripped = line.strip()
        if not stripped:
            continue
        try:
            item = json.loads(stripped)
        except json.JSONDecodeError as exc:
            raise DomainError(f"JSONL 第 {line_number} 行解析失败：{exc.msg}") from exc
        if not isinstance(item, dict):
            raise DomainError(f"JSONL 第 {line_number} 行不是对象，无法推断字段清单")
        total += 1
        if len(records) < sample_limit:
            records.append(item)
    if total == 0:
        raise DomainError("JSONL 文件没有有效行")
    return FileInspection(
        columns=_json_columns(records),
        row_count=total,
        sample_rows=tuple(records),
        statistics={"encoding": used_encoding, "root": "jsonl"},
    )


# ------------------------------------------------------------------ 二进制列存
def _inspect_parquet(data: bytes, sample_limit: int, encoding: str | None) -> FileInspection:
    del encoding
    import pyarrow.parquet as parquet

    reader = parquet.ParquetFile(io.BytesIO(data))
    schema = reader.schema_arrow
    columns = _columns((field.name, str(field.type)) for field in schema)
    # 只读**第一个 row group 的第一个 batch**：迭代器建立在 row group 边界上，
    # 读一个 batch 不会触发整文件解码。
    sample_rows: list[Mapping[str, Any]] = []
    for batch in reader.iter_batches(batch_size=sample_limit):
        sample_rows = [dict(row) for row in batch.to_pylist()]
        break
    return FileInspection(
        columns=columns,
        row_count=reader.metadata.num_rows,
        sample_rows=tuple(sample_rows),
        statistics={
            "row_groups": reader.metadata.num_row_groups,
            "compression": _parquet_compression(reader),
        },
    )


def _parquet_compression(reader: Any) -> str:
    """列出各列压缩算法（去重）；元数据拿不到时留空而不是编一个值。"""
    try:
        codecs = {
            reader.metadata.row_group(0).column(index).compression
            for index in range(reader.metadata.num_columns)
        }
    except (IndexError, AttributeError):
        return ""
    return ",".join(sorted(codecs))


def _inspect_orc(data: bytes, sample_limit: int, encoding: str | None) -> FileInspection:
    del encoding
    import pyorc

    stream = io.BytesIO(data)
    pairs: list[tuple[str, str]] = []
    sample_rows: list[Mapping[str, Any]] = []
    # pyorc.Reader 不是上下文管理器，只能手动清理。
    reader = pyorc.Reader(stream, struct_repr=pyorc.StructRepr.DICT)
    try:
        pairs = _parse_orc_schema(str(reader.schema))
        for index, row in enumerate(reader):
            if index >= sample_limit:
                break
            if not isinstance(row, Mapping):
                continue
            sample_rows.append({str(key): cell_to_plain(value) for key, value in row.items()})
    finally:
        close = getattr(reader, "close", None)
        if callable(close):
            close()
    if not pairs and sample_rows:
        pairs = [(key, "UNKNOWN") for key in sample_rows[0]]
    # ORC 读完后可回卷，但行数不在 Reader 的稳定 API 上，因此留空而不猜。
    return FileInspection(columns=_columns(pairs), row_count=None, sample_rows=tuple(sample_rows))


def _parse_orc_schema(schema: str) -> list[tuple[str, str]]:
    """解析 `struct<a:int,b:string>`；嵌套类型按尖括号配对切分。"""
    text = schema.strip()
    if not text.startswith("struct<") or not text.endswith(">"):
        return []
    body = text[len("struct<") : -1]
    pairs: list[tuple[str, str]] = []
    depth = 0
    current = ""
    for char in body:
        if char in "<(":
            depth += 1
        elif char in ">)":
            depth -= 1
        if char == "," and depth == 0:
            pairs.append(_split_orc_field(current))
            current = ""
            continue
        current += char
    if current:
        pairs.append(_split_orc_field(current))
    return [pair for pair in pairs if pair[0]]


def _split_orc_field(field: str) -> tuple[str, str]:
    name, separator, data_type = field.partition(":")
    return (name.strip().strip("`"), data_type.strip() if separator else "UNKNOWN")


# ------------------------------------------------------------------ Excel
def _inspect_xlsx(data: bytes, sample_limit: int, encoding: str | None) -> FileInspection:
    del encoding
    import openpyxl

    workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        sheet = workbook.worksheets[0]
        rows = sheet.iter_rows(values_only=True)
        header_row = next(rows, None)
        if header_row is None:
            raise DomainError("Excel 首个工作表为空")
        header = [
            str(cell).strip() if cell is not None and str(cell).strip() else f"column_{index}"
            for index, cell in enumerate(header_row, start=1)
        ]
        sample_rows: list[Mapping[str, Any]] = []
        for row in rows:
            if len(sample_rows) >= sample_limit:
                break
            sample_rows.append(
                {
                    header[index]: cell_to_plain(cell)
                    for index, cell in enumerate(row)
                    if index < len(header)
                }
            )
        return FileInspection(
            columns=_columns((name, "STRING") for name in header),
            row_count=(sheet.max_row - 1) if sheet.max_row else 0,
            sample_rows=tuple(sample_rows),
            statistics={"sheet": sheet.title, "sheets": list(workbook.sheetnames)},
        )
    finally:
        workbook.close()


# ------------------------------------------------------------------ XML
def _inspect_xml(data: bytes, sample_limit: int, encoding: str | None) -> FileInspection:
    del encoding
    try:
        root = SafeElementTree.fromstring(data)
    except SafeElementTree.ParseError as exc:
        raise DomainError(f"XML 解析失败：{exc}") from exc

    children = list(root)
    if not children:
        return _content_only("xml-text", root_tag=root.tag)
    # 重复出现最多的子标签即「记录」标签；只有一种子标签时也成立。
    tag_counts: dict[str, int] = {}
    for child in children:
        tag_counts[child.tag] = tag_counts.get(child.tag, 0) + 1
    record_tag = max(tag_counts, key=lambda tag: tag_counts[tag])
    records = [child for child in children if child.tag == record_tag]

    ordered: dict[str, str] = {}
    for record in records[:sample_limit]:
        for key in record.attrib:
            ordered.setdefault(f"@{key}", "STRING")
        for child in record:
            ordered.setdefault(child.tag, "STRING")
    if not ordered:
        # 记录元素里既没有子元素也没有属性：整段当文本处理，不产出样本。
        return _content_only("xml-text", root_tag=root.tag, record_tag=record_tag)

    sample_rows = []
    for record in records[:sample_limit]:
        row: dict[str, Any] = {
            f"@{key}": cell_to_plain(value) for key, value in record.attrib.items()
        }
        for child in record:
            row[child.tag] = cell_to_plain(child.text)
        sample_rows.append(row)
    return FileInspection(
        columns=_columns(ordered.items()),
        row_count=len(records),
        sample_rows=tuple(sample_rows),
        statistics={"root_tag": root.tag, "record_tag": record_tag, "child_tags": len(tag_counts)},
    )


# ------------------------------------------------------------------ Markdown
def _split_markdown_row(line: str) -> list[str]:
    cells = _MARKDOWN_CELL_SPLIT.split(line.strip())
    if cells and not cells[0].strip():
        cells = cells[1:]
    if cells and not cells[-1].strip():
        cells = cells[:-1]
    return [cell.strip() for cell in cells]


def _inspect_markdown(data: bytes, sample_limit: int, encoding: str | None) -> FileInspection:
    text, used_encoding = _decode(data, encoding)
    lines = text.splitlines()
    starts = _markdown_table_starts(lines)
    if not starts:
        return _content_only("markdown-text", tables=0, encoding=used_encoding)

    header = _split_markdown_row(lines[starts[0]])
    body: list[str] = []
    for candidate in lines[starts[0] + 2 :]:
        if not candidate.strip().startswith("|"):
            break
        body.append(candidate)
    return FileInspection(
        columns=_columns(
            (name or f"column_{position}", "STRING")
            for position, name in enumerate(header, start=1)
        ),
        row_count=len(body),
        sample_rows=tuple(
            {
                header[position]: cell_to_plain(cell)
                for position, cell in enumerate(_split_markdown_row(row))
                if position < len(header)
            }
            for row in body[:sample_limit]
        ),
        statistics={"encoding": used_encoding, "tables": len(starts)},
    )


def _markdown_table_starts(lines: Sequence[str]) -> list[int]:
    """找出所有 Markdown 表格的起始行号（表头行 + 下一行是对齐行）。"""
    starts: list[int] = []
    for index, line in enumerate(lines):
        if not line.strip().startswith("|"):
            continue
        if index + 1 < len(lines) and _MARKDOWN_ALIGNMENT_ROW.match(lines[index + 1]):
            starts.append(index)
    return starts


# ------------------------------------------------------------------ PDF
def _inspect_pdf(data: bytes, sample_limit: int, encoding: str | None) -> FileInspection:
    del sample_limit, encoding
    import pypdf

    try:
        reader = pypdf.PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise DomainError("PDF 已加密，不支持探查（不接受把文档口令当作连接凭据）")
        pages = len(reader.pages)
    except DomainError:
        raise
    except Exception as exc:  # pypdf 对损坏文件抛多种异常，统一收敛
        raise DomainError(f"PDF 解析失败：{type(exc).__name__}") from exc
    return _content_only("pdf-text", pages=pages)


_HANDLERS = {
    FileAssetFormat.CSV: _inspect_csv,
    FileAssetFormat.JSON: _inspect_json,
    FileAssetFormat.JSONL: _inspect_jsonl,
    FileAssetFormat.PARQUET: _inspect_parquet,
    FileAssetFormat.ORC: _inspect_orc,
    FileAssetFormat.XLSX: _inspect_xlsx,
    FileAssetFormat.XML: _inspect_xml,
    FileAssetFormat.MARKDOWN: _inspect_markdown,
    FileAssetFormat.PDF: _inspect_pdf,
}
